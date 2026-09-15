from pathlib import Path
from pptx import Presentation

ROOT = Path(r"c:\Users\user\Desktop\賞金獵人\line")
PPTX_PATH = ROOT / "01_企畫書" / "10頁投稿簡報" / "Temple_AI_OS_2026_LINE_AI_競賽簡報.pptx"

def inspect():
    prs = Presentation(PPTX_PATH)
    print(f"Total slides: {len(prs.slides)}")
    for i, slide in enumerate(prs.slides):
        shapes_info = []
        for s in slide.shapes:
            if s.has_text_frame:
                txt = s.text_frame.text.replace("\n", " ")[:40]
                shapes_info.append(f"Text({txt}...)")
            elif s.shape_type == 13: # Picture
                shapes_info.append("Picture")
            else:
                shapes_info.append(f"Shape({s.name})")
        print(f"\n--- Slide {i} ({len(slide.shapes)} shapes) ---")
        for s_desc in shapes_info[:6]:
            print(f"  {s_desc}")
        if len(shapes_info) > 6:
            print(f"  ... and {len(shapes_info) - 6} more shapes")

if __name__ == "__main__":
    inspect()
